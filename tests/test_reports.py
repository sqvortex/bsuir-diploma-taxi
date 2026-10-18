"""Тесты этапа 7: отчёты за период с выгрузкой в Excel (ФТ-18) и письма о брони (ФТ-7)."""

from datetime import date, datetime, time, timedelta
from io import BytesIO

import pytest
from django.core import mail
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from bookings.models import Booking
from bookings.notifications import bookings_to_remind, send_reminders
from bookings.services import create_booking
from reports.services import build_report, to_excel
from transport.models import Trip, Vehicle

pytestmark = pytest.mark.django_db


def make_booking(trip, from_stop, to_stop, seats, price, status, **extra):
    return Booking.objects.create(
        trip=trip, from_stop=from_stop, to_stop=to_stop, seats_count=seats, price=price,
        status=status, passenger_name="Тест", passenger_phone="+375291234567", **extra,
    )


@pytest.fixture
def past_trip(route, driver):
    """Рейс, который отправился вчера утром, и одна машина на 10 мест."""
    vehicle = Vehicle.objects.create(brand="Ford", plate_number="AB 0007-7", seats=10)
    yesterday = timezone.localdate() - timedelta(days=1)
    departure = timezone.make_aware(datetime.combine(yesterday, time(7, 30)))
    return Trip.objects.create(route=route, vehicle=vehicle, driver=driver, departure_at=departure,
                               status=Trip.Status.COMPLETED)


# ---------- Отчёты ----------

def test_report_numbers(past_trip, stops):
    minsk, bobruisk, gomel = stops
    make_booking(past_trip, minsk, gomel, 3, "105", Booking.Status.BOARDED)      # едет весь путь
    make_booking(past_trip, minsk, bobruisk, 2, "36", Booking.Status.BOARDED)    # только первый перегон
    make_booking(past_trip, bobruisk, gomel, 1, "17", Booking.Status.NO_SHOW)    # не явился
    make_booking(past_trip, minsk, gomel, 4, "140", Booking.Status.CANCELLED)    # отменена

    today = timezone.localdate()
    report = build_report(today - timedelta(days=3), today)
    summary = report.summary
    assert summary["trips"] == 1 and summary["bookings"] == 4
    assert summary["passengers"] == 5
    assert summary["revenue"] == 141.0               # неявка и отмена денег не приносят
    assert summary["no_show"] == 25.0 and summary["cancelled"] == 25.0
    # Самый занятый перегон: Минск – Бобруйск, там 3 + 2 + 0 = 5 мест из 10 (место не явившегося не считаем отдельно:
    # он ехал бы на втором перегоне, где занято 3 + 1 = 4).
    assert summary["load"] == 50.0
    row = report.by_trip.iloc[0]
    assert row["peak"] == 5 and row["load"] == 50.0
    assert report.by_route.iloc[0]["revenue"] == 141.0


def test_report_skips_future_cancelled_trips_and_other_routes(past_trip, trip, stops):
    minsk, _, gomel = stops
    make_booking(trip, minsk, gomel, 2, "70", Booking.Status.ACTIVE)   # рейс ещё не отправился
    cancelled = Trip.objects.create(route=past_trip.route, vehicle=past_trip.vehicle,
                                    departure_at=past_trip.departure_at - timedelta(hours=5),
                                    status=Trip.Status.CANCELLED)
    today = timezone.localdate()
    report = build_report(today - timedelta(days=3), today + timedelta(days=10))
    assert list(report.by_trip["trip"]) == [past_trip.pk]
    assert cancelled.pk not in list(report.by_trip["trip"])
    assert build_report(today - timedelta(days=3), today, route_id=past_trip.route_id + 100).summary["trips"] == 0


def test_empty_report(client):
    today = timezone.localdate()
    report = build_report(today - timedelta(days=5), today)
    assert report.summary["trips"] == 0 and report.summary["revenue"] == 0.0
    assert to_excel(report)  # пустой отчёт тоже выгружается


def test_excel_file(past_trip, stops):
    minsk, _, gomel = stops
    make_booking(past_trip, minsk, gomel, 2, "70", Booking.Status.BOARDED)
    today = timezone.localdate()
    workbook = load_workbook(BytesIO(to_excel(build_report(today - timedelta(days=2), today))))
    assert workbook.sheetnames == ["Итоги", "По маршрутам", "По дням", "По рейсам"]
    summary = {row[0].value: row[1].value for row in workbook["Итоги"].iter_rows(min_row=2)}
    assert summary["Выручка, BYN"] == 70.0 and summary["Рейсов"] == 1
    trips = list(workbook["По рейсам"].iter_rows(min_row=2, values_only=True))
    assert trips[0][1] == "Минск – Гомель" and trips[0][3] == 2


def test_reports_pages_and_access(client, past_trip, stops):
    minsk, _, gomel = stops
    make_booking(past_trip, minsk, gomel, 2, "70", Booking.Status.BOARDED)
    dispatcher = User.objects.create_user("+375292222222", "pass12345", role=User.Role.DISPATCHER)
    client.force_login(dispatcher)
    page = client.get(reverse("reports:overview"))
    assert page.status_code == 200
    text = page.content.decode()
    assert "Минск – Гомель" in text and "report-chart" in text and "Скачать Excel" in text

    response = client.get(reverse("reports:download"), {"from": "2020-01-01", "to": "2020-01-31"})
    assert response.status_code == 200
    assert response["Content-Type"].endswith("spreadsheetml.sheet")
    assert "report_20200101_20200131.xlsx" in response["Content-Disposition"]
    assert client.get(reverse("reports:overview"), {"from": "2020-01-01", "to": "2020-01-31"}).status_code == 200

    passenger = User.objects.create_user("+375293333333", "pass12345")
    client.force_login(passenger)
    assert client.get(reverse("reports:overview")).status_code == 403
    assert client.get(reverse("reports:download")).status_code == 403


# ---------- Письма ----------

@pytest.fixture
def passenger():
    return User.objects.create_user("+375293333333", "pass12345", first_name="Анна", email="anna@example.com")


def test_confirmation_email(trip, stops, passenger, django_capture_on_commit_callbacks):
    minsk, _, gomel = stops
    with django_capture_on_commit_callbacks(execute=True):
        booking = create_booking(trip=trip, from_stop=minsk, to_stop=gomel, seats_count=2,
                                 passenger_name="Анна", passenger_phone=passenger.phone,
                                 passenger=passenger, created_by=passenger)
    [message] = mail.outbox
    assert message.to == ["anna@example.com"]
    assert booking.code in message.subject
    assert "Минск – Гомель" in message.body and "Мест: 2" in message.body
    assert reverse("bookings:detail", args=[booking.pk]) in message.body


def test_no_email_without_address(trip, stops, django_capture_on_commit_callbacks):
    minsk, _, gomel = stops
    nobody = User.objects.create_user("+375294444444", "pass12345")  # почта не указана
    with django_capture_on_commit_callbacks(execute=True):
        create_booking(trip=trip, from_stop=minsk, to_stop=gomel, seats_count=1,
                       passenger_name="Без почты", passenger_phone=nobody.phone, passenger=nobody)
        create_booking(trip=trip, from_stop=minsk, to_stop=gomel, seats_count=1,
                       passenger_name="Гость", passenger_phone="+375295555555")  # бронь по звонку
    assert mail.outbox == []


def test_mail_error_does_not_break_booking(trip, stops, passenger, settings, django_capture_on_commit_callbacks):
    settings.EMAIL_BACKEND = "tests.test_reports.BrokenBackend"
    minsk, _, gomel = stops
    with django_capture_on_commit_callbacks(execute=True):
        booking = create_booking(trip=trip, from_stop=minsk, to_stop=gomel, seats_count=1,
                                 passenger_name="Анна", passenger_phone=passenger.phone, passenger=passenger)
    assert Booking.objects.filter(pk=booking.pk).exists()


class BrokenBackend:
    def __init__(self, *args, **kwargs):
        pass

    def send_messages(self, messages):
        raise ConnectionError("почтовый сервер недоступен")


def tomorrow_trip(route, driver, template_time=time(9, 0)):
    vehicle = Vehicle.objects.create(brand="Ford", plate_number="AB 0009-7", seats=10)
    tomorrow = timezone.localdate() + timedelta(days=1)
    departure = timezone.make_aware(datetime.combine(tomorrow, template_time))
    return Trip.objects.create(route=route, vehicle=vehicle, driver=driver, departure_at=departure)


def test_reminders_sent_once(route, stops, driver, passenger, trip):
    minsk, _, gomel = stops
    soon = tomorrow_trip(route, driver)
    mine = make_booking(soon, minsk, gomel, 1, "35", Booking.Status.ACTIVE, passenger=passenger)
    make_booking(soon, minsk, gomel, 1, "35", Booking.Status.CANCELLED, passenger=passenger)   # отменена
    make_booking(trip, minsk, gomel, 1, "35", Booking.Status.ACTIVE, passenger=passenger)      # через неделю
    make_booking(soon, minsk, gomel, 1, "35", Booking.Status.ACTIVE)                           # без аккаунта

    assert bookings_to_remind() == [mine]
    assert send_reminders() == 1
    [message] = mail.outbox
    assert "завтра" in message.body and mine.code in message.subject
    mine.refresh_from_db()
    assert mine.reminded_at is not None
    assert send_reminders() == 0   # повторный запуск никому не пишет
    call_command("send_reminders", stdout=None)
    assert len(mail.outbox) == 1
